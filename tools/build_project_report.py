from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "report_artifacts"
REPORT_PATH = ROOT / "BaoCao_MultiAgent_ChiTiet.docx"
DIAGRAM_PATH = OUT_DIR / "multi_agent_pipeline.png"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "17365D"
LIGHT_BLUE = "DCE6F1"
LIGHTER_BLUE = "EEF5FB"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D9E2F3"
TEXT = "1F2937"
MUTED = "5B6573"
GREEN = "2E7D5B"
AMBER = "9A6700"
WHITE = "FFFFFF"

NUMBERING_IDS = {}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_table_grid(table, widths: Sequence[int], indent_twips: int = 120) -> None:
    table.autofit = False
    tbl = table._tbl
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_twips))
    tbl_ind.set(qn("w:type"), "dxa")
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(idx, len(widths) - 1)])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_borders(table, color="D6DCE4", size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_run_font(run, name="Calibri", size=11, color=TEXT, bold=False, italic=False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_text(paragraph, text: str, *, bold=False, italic=False, color=TEXT, size=11, font="Calibri"):
    run = paragraph.add_run(text)
    set_run_font(run, font, size, color, bold, italic)
    return run


def add_para(doc, text="", *, style=None, align=None, before=0, after=6, line=1.10, keep=False):
    p = doc.add_paragraph(style=style)
    if text:
        add_text(p, text)
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep
    if align is not None:
        p.alignment = align
    return p


def ensure_numbering(doc, kind: str) -> int:
    key = (id(doc), kind)
    if key in NUMBERING_IDS:
        return NUMBERING_IDS[key]
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "multilevel")
    abstract.append(multi)
    for level in (0, 1):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), "•" if kind == "bullet" else f"%{level + 1}.")
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), str(720 + level * 360))
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(720 + level * 360))
        ind.set(qn("w:hanging"), "360")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "160")
        spacing.set(qn("w:line"), "280")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.extend([tabs, ind, spacing])
        lvl.extend([start, num_fmt, lvl_text, lvl_jc, p_pr])
        if kind == "bullet":
            r_pr = OxmlElement("w:rPr")
            r_fonts = OxmlElement("w:rFonts")
            r_fonts.set(qn("w:ascii"), "Symbol")
            r_fonts.set(qn("w:hAnsi"), "Symbol")
            r_pr.append(r_fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    NUMBERING_IDS[key] = num_id
    return num_id


def apply_numbering(paragraph, num_id: int, level: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])


def add_bullet(doc, text: str, level=0):
    p = doc.add_paragraph()
    apply_numbering(p, ensure_numbering(doc, "bullet"), level)
    add_text(p, text)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    return p


def add_number(doc, text: str, level=0):
    p = doc.add_paragraph()
    apply_numbering(p, ensure_numbering(doc, "decimal"), level)
    add_text(p, text)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    return p


def paragraph_rule(paragraph, color=BLUE, size="12"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def add_heading(doc, text: str, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    add_text(
        p,
        text,
        bold=True,
        color=BLUE if level <= 2 else DARK_BLUE,
        size={1: 16, 2: 13, 3: 12}.get(level, 11),
    )
    return p


def add_callout(doc, title: str, body: str, *, tone="blue"):
    palette = {
        "blue": (LIGHTER_BLUE, BLUE),
        "green": ("EAF5EF", GREEN),
        "amber": ("FFF4D6", AMBER),
    }
    fill, accent = palette[tone]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_grid(table, [140, 9220])
    set_table_borders(table, color=fill, size="0")
    left, right = table.rows[0].cells
    set_cell_shading(left, accent)
    set_cell_shading(right, fill)
    set_cell_margins(left, 80, 120, 80, 120)
    set_cell_margins(right, 130, 120, 130, 120)
    p = right.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    add_text(p, title, bold=True, color=accent, size=11)
    p2 = right.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    add_text(p2, body, color=TEXT, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers: Sequence[str], rows: Iterable[Sequence[str]], widths: Sequence[int], font_size=9.2):
    rows = list(rows)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_grid(table, widths)
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_text(p, text, bold=True, color=NAVY, size=font_size)
    for row_values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            if len(table.rows) % 2 == 1:
                set_cell_shading(cell, "FAFBFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_text(p, str(value), color=TEXT, size=font_size)
    return table


def add_code_block(doc, text: str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_grid(table, [9360], indent_twips=180)
    set_table_borders(table, color="DCE2E8", size="6")
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F6F8FA")
    set_cell_margins(cell, 150, 180, 150, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for i, line in enumerate(text.splitlines()):
        if i:
            p.add_run().add_break()
        add_text(p, line, font="Consolas", size=8.7, color="27364B")
    return table


def page_break(doc):
    doc.add_page_break()


def draw_arrow(draw, start, end, color="#52769B", width=5):
    draw.line([start, end], fill=color, width=width)
    x2, y2 = end
    x1, y1 = start
    if abs(x2 - x1) > abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        pts = [(x2, y2), (x2 - direction * 18, y2 - 11), (x2 - direction * 18, y2 + 11)]
    else:
        direction = 1 if y2 > y1 else -1
        pts = [(x2, y2), (x2 - 11, y2 - direction * 18), (x2 + 11, y2 - direction * 18)]
    draw.polygon(pts, fill=color)


def make_diagram(path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", (1500, 950), "#F7FAFD")
    draw = ImageDraw.Draw(img)
    font_path = Path("C:/Windows/Fonts/arial.ttf")
    bold_path = Path("C:/Windows/Fonts/arialbd.ttf")
    font = ImageFont.truetype(str(font_path), 25)
    small = ImageFont.truetype(str(font_path), 21)
    bold = ImageFont.truetype(str(bold_path), 27)

    def box(x, y, w, h, title, subtitle="", fill="#FFFFFF", outline="#2E74B5"):
        draw.rounded_rectangle((x, y, x + w, y + h), radius=18, fill=fill, outline=outline, width=4)
        bbox = draw.textbbox((0, 0), title, font=bold)
        draw.text((x + (w - (bbox[2] - bbox[0])) / 2, y + 20), title, font=bold, fill="#17365D")
        if subtitle:
            lines = subtitle.split("\n")
            yy = y + 61
            for line in lines:
                bb = draw.textbbox((0, 0), line, font=small)
                draw.text((x + (w - (bb[2] - bb[0])) / 2, yy), line, font=small, fill="#4B6074")
                yy += 28

    box(50, 48, 250, 110, "Input JSON", "EC_001 ... EC_050", fill="#EEF5FB")
    box(390, 48, 300, 110, "Coordinator", "điều phối & tổng hợp", fill="#DCE6F1")
    box(790, 48, 300, 110, "DataLoader", "lọc CSV theo order_id", fill="#EAF5EF", outline="#2E7D5B")
    draw_arrow(draw, (300, 103), (390, 103))
    draw_arrow(draw, (690, 103), (790, 103))

    box(75, 300, 370, 135, "Order & Seller", "trạng thái · item · seller\nshipping limit")
    box(565, 300, 370, 135, "Payment", "đối soát item + freight\nsplit payment")
    box(1055, 300, 370, 135, "Delivery", "giao đúng/trễ\nseller hay logistics")
    draw.line([(940, 158), (940, 235)], fill="#52769B", width=5)
    draw.line([(260, 235), (1240, 235)], fill="#52769B", width=5)
    draw_arrow(draw, (260, 235), (260, 300))
    draw_arrow(draw, (750, 235), (750, 300))
    draw_arrow(draw, (1240, 235), (1240, 300))

    box(565, 535, 370, 125, "Policy", "áp dụng EC_POLICY_V1\nissue · cause · refund · action", fill="#FFF4D6", outline="#9A6700")
    draw.line([(260, 435), (260, 485), (750, 485)], fill="#52769B", width=5)
    draw.line([(750, 435), (750, 535)], fill="#52769B", width=5)
    draw.line([(1240, 435), (1240, 485), (750, 485)], fill="#52769B", width=5)
    draw_arrow(draw, (750, 485), (750, 535))

    box(85, 755, 300, 120, "Evidence", "chọn evidence IDs", fill="#EEF5FB")
    box(475, 755, 300, 120, "Verifier", "chuẩn hóa & validate", fill="#EAF5EF", outline="#2E7D5B")
    box(865, 755, 300, 120, "Output JSON", "50 kết quả hợp lệ", fill="#DCE6F1")
    box(1215, 755, 230, 120, "Trace", "audit JSONL", fill="#F2F4F7", outline="#68778A")
    draw.line([(750, 660), (750, 705), (235, 705)], fill="#52769B", width=5)
    draw_arrow(draw, (235, 705), (235, 755))
    draw_arrow(draw, (385, 815), (475, 815))
    draw_arrow(draw, (775, 815), (865, 815))
    draw_arrow(draw, (1165, 815), (1215, 815))

    img.save(path, quality=95)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        style = doc.styles[list_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.10


def add_header_footer(doc: Document):
    for section in doc.sections:
        header = section.header
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_text(p, "K3 DAY 09  ·  MULTI-AGENT A2A", bold=True, color="718096", size=8.5)

        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, "Báo cáo thiết kế hệ thống  |  ", color="718096", size=8.5)
        run = p.add_run()
        fld_char1 = OxmlElement("w:fldChar")
        fld_char1.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        fld_char2 = OxmlElement("w:fldChar")
        fld_char2.set(qn("w:fldCharType"), "end")
        run._r.append(fld_char1)
        run._r.append(instr)
        run._r.append(fld_char2)
        set_run_font(run, size=8.5, color="718096")


def build_report():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    make_diagram(DIAGRAM_PATH)
    doc = Document()
    configure_document(doc)

    # Cover page — editorial cover pattern.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(55)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "BÁO CÁO KỸ THUẬT", bold=True, color=BLUE, size=13)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(12)
    add_text(p, "THIẾT KẾ VÀ LUỒNG HOẠT ĐỘNG\nHỆ THỐNG MULTI-AGENT", bold=True, color=NAVY, size=25)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(26)
    add_text(p, "Olist E-commerce Dispute Resolution — Bản phân tích chi tiết 7 agent", color=MUTED, size=14)

    divider = doc.add_paragraph()
    divider.alignment = WD_ALIGN_PARAGRAPH.CENTER
    divider.paragraph_format.left_indent = Inches(2.45)
    divider.paragraph_format.right_indent = Inches(2.45)
    divider.paragraph_format.space_after = Pt(0)
    paragraph_rule(divider, color=BLUE, size="18")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(36)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "Môn học / Bài lab", color=MUTED, size=9.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "K3 Day 09 — Multi-Agent A2A", bold=True, color=TEXT, size=12)

    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_grid(meta, [2600, 4200])
    set_table_borders(meta, color="FFFFFF", size="0")
    values = [
        ("Người thực hiện", "Vũ Việt Anh"),
        ("MSSV", "2A202601107"),
        ("Phiên bản policy", "EC_POLICY_V1"),
        ("Ngày báo cáo", "05/08/2026"),
    ]
    for row, (label, value) in zip(meta.rows, values):
        set_cell_margins(row.cells[0], 80, 120, 80, 120)
        set_cell_margins(row.cells[1], 80, 120, 80, 120)
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_text(row.cells[0].paragraphs[0], label, color=MUTED, size=10)
        add_text(row.cells[1].paragraphs[0], value, bold=True, color=TEXT, size=10)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(70)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "Python · Pandas · Pydantic · OpenAI SDK", color="718096", size=9.5)

    page_break(doc)

    # Page 2
    add_heading(doc, "1. Tóm tắt dự án", 1)
    add_para(
        doc,
        "Dự án xây dựng một hệ thống multi-agent để điều tra 50 khiếu nại thương mại điện tử trên bộ dữ liệu Olist. "
        "Mỗi case được truy xuất theo claimed_order_id, đối chiếu trạng thái đơn, item, seller, thanh toán và các mốc giao hàng; "
        "sau đó hệ thống xác định vấn đề chính, nguyên nhân gốc, bên chịu trách nhiệm, bằng chứng, số tiền hoàn và hành động xử lý.",
        after=8,
    )
    add_callout(
        doc,
        "Kết quả chính",
        "Pipeline xử lý đủ 50 input JSON, tạo 50 output JSON theo schema và ghi trace của lượt chạy mới nhất. "
        "Thiết kế tách nghiệp vụ thành 7 agent có contract handoff rõ ràng.",
        tone="green",
    )
    add_heading(doc, "1.1. Mục tiêu thiết kế", 2)
    for text in (
        "Phân chia trách nhiệm theo domain để mỗi agent chỉ xử lý một nhóm dữ liệu hoặc một bước quyết định.",
        "Ưu tiên dữ liệu CSV có thể kiểm chứng; không tự tạo sự kiện, ID hoặc số tiền không tồn tại.",
        "Áp dụng EC_POLICY_V1 đúng thứ tự ưu tiên và tạo kết quả có tính tái lập.",
        "Tách affected entities khỏi evidence để vừa báo cáo đầy đủ đối tượng liên quan, vừa tránh evidence dư thừa.",
        "Có lớp kiểm định cuối cùng và trace để audit từng handoff.",
    ):
        add_bullet(doc, text)
    add_heading(doc, "1.2. Phạm vi dữ liệu", 2)
    add_para(doc, "DataLoader nạp 6 bảng cần cho luồng hiện tại: orders, order_items, order_payments, customers, sellers và products. Ba bảng được truy vấn trực tiếp cho từng case là orders, order_items và order_payments.")
    add_table(
        doc,
        ["Khóa liên kết", "Ý nghĩa trong pipeline"],
        [
            ("orders.order_id → order_items.order_id", "Lấy item, seller, giá, freight và shipping limit."),
            ("orders.order_id → order_payments.order_id", "Lấy các payment row và tổng số tiền khách đã trả."),
            ("order_items.seller_id → sellers.seller_id", "Định danh seller liên quan hoặc seller chịu trách nhiệm."),
        ],
        [3500, 5860],
        font_size=9.5,
    )

    page_break(doc)

    # Page 3
    add_heading(doc, "2. Kiến trúc tổng thể", 1)
    add_para(
        doc,
        "Kiến trúc sử dụng Coordinator làm điểm vào duy nhất. Dữ liệu thô được nạp một lần, ba agent chuyên môn phân tích các góc nhìn độc lập, "
        "Policy hợp nhất kết quả, Evidence chọn bằng chứng và Verifier kiểm tra output trước khi ghi file.",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    picture = p.add_run().add_picture(str(DIAGRAM_PATH), width=Inches(6.5))
    picture._inline.docPr.set(
        "descr",
        "Sơ đồ pipeline: Input JSON qua Coordinator và DataLoader, ba specialist Order-Seller, Payment, Delivery, sau đó Policy, Evidence, Verifier, Output JSON và Trace.",
    )
    picture._inline.docPr.set("title", "Luồng xử lý hệ thống multi-agent Olist")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    add_text(p, "Hình 1. Luồng xử lý end-to-end và các điểm handoff", italic=True, color=MUTED, size=9)
    add_callout(
        doc,
        "Đặc điểm quan trọng",
        "Đây là multi-agent theo cách phân rã trách nhiệm bằng các lớp Python. Ba specialist hiện được gọi tuần tự trong Coordinator; "
        "chúng độc lập về logic nhưng chưa chạy song song.",
        tone="blue",
    )

    page_break(doc)

    # Page 4
    add_heading(doc, "3. Nhiệm vụ của 7 agent", 1)
    agent_rows = [
        ("1", "CoordinatorAgent", "Input case + DataLoader", "Trích order_id; gọi 6 agent; tổng hợp candidate output.", "Final output + trace handoff"),
        ("2", "OrderSellerAgent", "Order + item rows", "Chuẩn hóa trạng thái, item, seller và shipping limit; sắp xếp ID ổn định.", "order_seller result"),
        ("3", "PaymentAgent", "Payment + item rows", "Tính item/freight/payment; kiểm tra sai số 0,10 BRL; nhận diện split payment.", "payment result"),
        ("4", "DeliveryAgent", "Timestamp + shipping limit", "Xác định giao trễ; phân biệt seller handoff trễ hay logistics giao trễ.", "delivery result"),
        ("5", "PolicyAgent", "3 specialist results", "Áp dụng 6 luật EC_POLICY_V1 theo thứ tự; quyết định issue, cause, party, refund, action.", "policy result"),
        ("6", "EvidenceAgent", "Order/payment/policy result", "Chọn evidence có thật, có thứ tự; thêm seller khi seller chịu trách nhiệm.", "evidence_ids"),
        ("7", "VerifierAgent", "Candidate output", "Cắt giới hạn mảng, làm tròn tiền, đồng bộ status/refund, chặn confidence, validate Pydantic.", "Output hợp lệ hoặc lỗi"),
    ]
    add_table(doc, ["#", "Agent", "Input", "Xử lý chính", "Handoff"], agent_rows, [420, 1700, 1700, 3600, 1940], font_size=8.35)
    add_heading(doc, "3.1. Ranh giới quyền truy cập", 2)
    add_para(doc, "Coordinator thông qua DataLoader là nơi lấy dữ liệu CSV. Specialist chỉ nhận context đã lọc; Policy không đọc CSV; Evidence không tự suy diễn; Verifier không tạo quyết định nghiệp vụ mới. Chỉ main.py ghi output/*.json và Tracer ghi logging/trace.jsonl.")
    add_callout(
        doc,
        "Vì sao cách chia này hợp lý?",
        "Mỗi agent có một nhiệm vụ dễ kiểm thử, giảm việc trộn lẫn logic tài chính với logic giao hàng, đồng thời tạo điểm kiểm soát rõ trước khi xuất kết quả.",
        tone="green",
    )

    page_break(doc)

    # Detailed agent profiles — expanded edition.
    add_heading(doc, "3.2. CoordinatorAgent — điều phối toàn bộ pipeline", 1)
    add_para(doc, "Mã nguồn: agents/coordinator_agent.py · Lớp: CoordinatorAgent · Phương thức chính: process(input_case)", after=8)
    add_callout(doc, "Vai trò", "Coordinator là bộ điều khiển luồng. Agent này không tự phân tích mọi domain; nó chuẩn bị context, giao việc, thu kết quả, dựng output ứng viên và chỉ trả kết quả sau khi Verifier chấp nhận.", tone="blue")
    add_heading(doc, "Input và output", 2)
    add_table(
        doc,
        ["Thành phần", "Nội dung chi tiết"],
        [
            ("Input", "Một case JSON gồm case_id, customer_request.claimed_order_id và policy_version."),
            ("Dependency", "DataLoader, LLMClient và 6 sub-agent được khởi tạo trong constructor."),
            ("Output", "Tuple (final_output, agent_traces). final_output đã qua Verifier; traces chứa output của từng sub-agent."),
            ("Lỗi", "Raise ValueError nếu Verifier trả valid=false hoặc nếu Policy không tìm thấy rule phù hợp."),
        ],
        [2200, 7160],
        font_size=9.3,
    )
    add_heading(doc, "Các bước xử lý bên trong", 2)
    for text in (
        "Trích case_id và claimed_order_id; đây là khóa xuyên suốt của một lượt điều tra.",
        "Gọi DataLoader.get_order_context() đúng một lần để lấy order_info, items và payments.",
        "Gọi lần lượt OrderSellerAgent, PaymentAgent và DeliveryAgent; lưu từng kết quả vào agent_traces.",
        "Tạo policy_context chỉ gồm ba kết quả chuyên môn rồi gọi PolicyAgent.",
        "Dựng ID chuẩn cho item theo <order_id>:<order_item_id> và payment theo <order_id>:<payment_sequential>.",
        "Gọi EvidenceAgent, ghép assessment, entities, cause, financials và actions thành final_dict.",
        "Xử lý riêng order không có item: xóa item/seller ID và đặt item/freight total bằng 0.",
        "Gọi VerifierAgent; chỉ trả dữ liệu khi valid=true.",
    ):
        add_bullet(doc, text)
    add_callout(doc, "Điểm thiết kế", "Coordinator là nơi duy nhất biết toàn bộ đồ thị agent. Nhờ đó các specialist không phụ thuộc lẫn nhau và Policy chỉ làm việc với contract đã chuẩn hóa.", tone="green")

    page_break(doc)

    add_heading(doc, "3.3. OrderSellerAgent — phân tích đơn hàng và seller", 1)
    add_para(doc, "Mã nguồn: agents/order_seller_agent.py · Lớp: OrderSellerAgent · Phương thức: process(context)", after=8)
    add_heading(doc, "Mục đích nghiệp vụ", 2)
    add_para(doc, "Agent này biến dữ liệu order/item thô thành một góc nhìn ổn định cho các bước phía sau. Nó trả trạng thái đơn, danh sách item, seller liên quan và mốc shipping_limit_date nhưng không kết luận bên chịu trách nhiệm hoặc số tiền hoàn.")
    add_heading(doc, "Logic xử lý", 2)
    for text in (
        "Đọc order_info; nếu order không tồn tại thì các trường order_status và order_id nhận None.",
        "Đặt has_items dựa trên số lượng item row.",
        "Sắp xếp item theo order_item_id để output không phụ thuộc thứ tự dòng CSV.",
        "Chuẩn hóa mỗi item thành order_item_id, product_id, seller_id, shipping_limit_date, price và freight_value.",
        "Thu seller_id bằng set để loại trùng rồi sorted để giữ thứ tự xác định.",
        "Truyền order_delivered_carrier_date sang output để hỗ trợ phân tích bàn giao.",
    ):
        add_bullet(doc, text)
    add_heading(doc, "Contract bàn giao", 2)
    add_code_block(
        doc,
        "{\n"
        "  order_status, order_id, has_items,\n"
        "  items: [{order_item_id, product_id, seller_id,\n"
        "           shipping_limit_date, price, freight_value}],\n"
        "  seller_ids, order_delivered_carrier_date\n"
        "}"
    )
    add_callout(doc, "Tình huống biên", "Order unavailable có thể không có item row. Khi đó agent trả items=[] và seller_ids=[]; Coordinator tiếp tục bảo đảm các tổng item/freight bằng 0 trong final output.", tone="amber")

    page_break(doc)

    add_heading(doc, "3.4. PaymentAgent — đối soát tài chính", 1)
    add_para(doc, "Mã nguồn: agents/payment_agent.py · Lớp: PaymentAgent · Phương thức: process(context)", after=8)
    add_heading(doc, "Các phép tính cốt lõi", 2)
    add_table(
        doc,
        ["Biến", "Công thức / điều kiện", "Ý nghĩa"],
        [
            ("total_payment_brl", "Σ payment_value", "Tổng khách đã thanh toán trên mọi payment row."),
            ("total_item_brl", "Σ item.price", "Tổng giá sản phẩm trong order."),
            ("total_freight_brl", "Σ item.freight_value", "Tổng phí vận chuyển."),
            ("expected_total", "item total + freight total", "Giá trị dự kiến từ item rows."),
            ("payment_matches_order", "|payment − expected| ≤ 0,10", "Cho phép sai số tối đa 0,10 BRL."),
            ("has_split_payment", "số payment row ≥ 2", "Nhận diện giao dịch thanh toán chia nhỏ."),
        ],
        [2200, 3000, 4160],
        font_size=9.0,
    )
    add_heading(doc, "Quy trình và handoff", 2)
    for text in (
        "Đọc payments và items từ cùng data_context, không đọc lại CSV.",
        "Chuyển payment_value, price và freight_value sang float trước khi cộng.",
        "Sắp xếp payment theo payment_sequential và chỉ bàn giao hai trường cần cho evidence: sequence và value.",
        "Trả cả tổng tiền và các cờ đối soát để Policy không phải lặp lại phép tính.",
    ):
        add_bullet(doc, text)
    add_callout(doc, "Lưu ý", "payment_value là giá trị của từng payment row, không phải giá trị mỗi installment. Vì vậy hệ thống cộng các row, không nhân với payment_installments.", tone="blue")

    add_heading(doc, "3.5. DeliveryAgent — xác định giao trễ và nguyên nhân", 1)
    add_para(doc, "Mã nguồn: agents/delivery_agent.py · Lớp: DeliveryAgent · Phương thức: process(context)", after=6)
    add_para(doc, "Agent so sánh ngày khách nhận với ngày dự kiến. Nếu giao trễ, nó tiếp tục so sánh ngày carrier nhận hàng với shipping limit của từng item để phân biệt lỗi seller và logistics.")
    for text in (
        "is_late_delivery=true khi delivered_customer_date và estimated_delivery_date cùng tồn tại, đồng thời actual > estimated.",
        "carrier_pickup_late=true cho một item khi delivered_carrier_date > shipping_limit_date.",
        "Nếu đơn trễ và có ít nhất một item carrier nhận muộn: late_cause=seller.",
        "Nếu đơn trễ nhưng không có item bàn giao muộn: late_cause=logistics.",
        "Nếu đơn không trễ: late_cause=none.",
    ):
        add_bullet(doc, text)
    add_callout(doc, "Giới hạn hiện tại", "Nếu chưa có ngày giao thực tế, code không tự coi đơn là trễ dù estimated date đã qua. Nhánh này đang để pass để tránh suy diễn ngoài dữ liệu chính thức.", tone="amber")

    page_break(doc)

    add_heading(doc, "3.6. PolicyAgent — bộ máy quyết định EC_POLICY_V1", 1)
    add_para(doc, "Mã nguồn: agents/policy_agent.py · Lớp: PolicyAgent · Phương thức: process(context)", after=8)
    add_callout(doc, "Trách nhiệm", "PolicyAgent biến fact đã chuẩn hóa thành quyết định nghiệp vụ: primary issue, root cause, responsible party, refund, action, case status và confidence.", tone="blue")
    add_heading(doc, "Input contract", 2)
    add_para(doc, "context gồm ba nhánh order_seller, payment và delivery. PolicyAgent không truy cập DataLoader hoặc CSV; điều này ngăn lớp quyết định thay đổi dữ kiện nguồn.")
    add_heading(doc, "Cơ chế ưu tiên luật", 2)
    for text in (
        "Ưu tiên 1–2: canceled/unavailable nhưng đã trả tiền → platform chịu trách nhiệm, hoàn toàn bộ payment.",
        "Ưu tiên 3–4: giao trễ → xác định seller hoặc logistics, hoàn toàn bộ freight.",
        "Ưu tiên 5: có split payment và tổng tiền khớp → không hoàn, giải thích giao dịch hợp lệ.",
        "Ưu tiên 6: claim giao trễ không được dữ liệu hỗ trợ và payment khớp → không hoàn, bác yêu cầu.",
        "Không có luật phù hợp → raise ValueError thay vì tạo kết luận mặc định.",
    ):
        add_bullet(doc, text)
    add_heading(doc, "Output contract", 2)
    add_code_block(
        doc,
        "{\n"
        "  primary_issue, root_cause_code,\n"
        "  responsible_party_type, responsible_party_id,\n"
        "  recommended_refund_brl, resolution_actions,\n"
        "  case_status, confidence\n"
        "}"
    )
    add_heading(doc, "Case status và confidence", 2)
    add_para(doc, "case_status được đặt action_required khi refund > 0, ngược lại no_action. confidence luôn bằng 1.0 khi case khớp trọn vẹn một rule. Đây là độ chắc chắn quy ước của luật xác định, không phải xác suất thống kê hoặc output của LLM.")
    add_callout(doc, "Quan hệ với LLMClient", "Constructor nhận LLMClient nhưng process() không gọi chat() hay get_json(). Quyết định thực tế hiện được tạo hoàn toàn bằng if/elif trên fact từ specialist.", tone="amber")

    page_break(doc)

    add_heading(doc, "3.7. EvidenceAgent — chọn bằng chứng có thể kiểm chứng", 1)
    add_para(doc, "Mã nguồn: agents/evidence_agent.py · Lớp: EvidenceAgent · Phương thức: process(context)", after=8)
    add_heading(doc, "Mục tiêu", 2)
    add_para(doc, "EvidenceAgent không tìm nguyên nhân mới. Nó nhận quyết định Policy cùng output order/payment rồi dựng tập evidence ID tối thiểu nhưng đủ chứng minh kết luận và các số tiền đã báo cáo.")
    add_heading(doc, "Quy tắc chọn evidence", 2)
    for text in (
        "Luôn thêm order:<order_id> để neo case vào bản ghi đơn hàng.",
        "Thêm item:<order_id>:<order_item_id> cho mọi item hiện có vì item rows tạo ra item_total và freight_total.",
        "Thêm payment:<order_id>:<payment_sequential> cho mọi payment row vì chúng tạo ra payment_total và tham gia điều kiện refund/reconciliation.",
        "Chỉ thêm seller:<seller_id> nếu responsible_party_type là seller; seller chỉ liên quan nhưng không chịu trách nhiệm sẽ không thành evidence.",
        "Thêm policy:<root_cause_code> để chỉ rõ rule nghiệp vụ đã dùng.",
    ):
        add_bullet(doc, text)
    add_callout(doc, "Phân biệt quan trọng", "affected_entities trả lời “những đối tượng nào liên quan?”, còn evidence_ids trả lời “những bản ghi nào chứng minh kết luận?”. Hai danh sách không bắt buộc giống nhau.", tone="green")

    add_heading(doc, "3.8. VerifierAgent — cổng chất lượng cuối", 1)
    add_para(doc, "Mã nguồn: agents/verifier_agent.py · Lớp: VerifierAgent · Phương thức: process(data)", after=6)
    add_heading(doc, "Năm lớp kiểm tra", 2)
    for text in (
        "Giới hạn độ dài: tối đa 5 ID mỗi entity set, 10 evidence, 3 causes, 3 parties và 5 actions.",
        "Làm tròn bốn trường tài chính về 2 chữ số thập phân.",
        "Sửa case_status theo recommended_refund_brl để tránh mâu thuẫn logic.",
        "Chặn confidence trong miền [0,1].",
        "Khởi tạo OutputSchema bằng Pydantic; trả valid=false và error nếu schema không hợp lệ.",
    ):
        add_bullet(doc, text)
    add_heading(doc, "Điều Verifier không làm", 2)
    add_para(doc, "Verifier không chọn lại policy, không đổi responsible party và chưa tra ngược evidence ID vào CSV. Nó là validator cấu trúc/constraint và normalization layer, không phải một policy engine thứ hai.")
    add_callout(doc, "Handoff cuối", "valid=true → Coordinator trả data cho main.py ghi JSON. valid=false → Coordinator raise ValueError, case không được coi là thành công.", tone="blue")

    page_break(doc)

    # Page 5
    add_heading(doc, "4. Luồng xử lý chi tiết", 1)
    steps = [
        "main.py khởi tạo DataLoader, LLMClient, Tracer và CoordinatorAgent. Tracer xóa trace cũ để chỉ lưu lượt chạy mới nhất.",
        "main.py lấy danh sách input/EC_*.json theo thứ tự và đọc từng case.",
        "Coordinator trích case_id và customer_request.claimed_order_id.",
        "DataLoader lọc ba bảng orders, order_items và order_payments theo order_id, chuyển NaN thành None rồi trả data_context.",
        "OrderSellerAgent, PaymentAgent và DeliveryAgent xử lý cùng data_context để tạo ba kết quả chuyên môn.",
        "PolicyAgent nhận policy_context gồm đúng ba kết quả trên và chọn luật EC_POLICY_V1 đầu tiên thỏa điều kiện.",
        "Coordinator dựng affected_entities, root_cause_analysis và financial_resolution; EvidenceAgent dựng evidence_ids.",
        "VerifierAgent chuẩn hóa, kiểm tra constraint và validate OutputSchema. Nếu không hợp lệ, case bị báo lỗi; nếu hợp lệ, dữ liệu được ghi ra output/<case_id>.json.",
        "Tracer ghi agent_traces và final_output thành một dòng JSONL phục vụ audit.",
    ]
    for step in steps:
        add_number(doc, step)
    add_heading(doc, "4.1. Contract handoff", 2)
    add_table(
        doc,
        ["Từ → đến", "Contract chính", "Không được làm"],
        [
            ("DataLoader → Specialists", "order_info, items[], payments[]", "Không trả dữ liệu ngoài claimed_order_id."),
            ("Specialists → Policy", "fact đã chuẩn hóa: totals, late cause, status", "Không quyết định refund thay Policy."),
            ("Policy → Coordinator/Evidence", "issue, cause, party, refund, actions, confidence", "Không dựng evidence ID không tồn tại."),
            ("Coordinator → Verifier", "candidate output đủ 7 nhóm trường", "Không bỏ qua schema/constraint."),
            ("Verifier → main.py", "{valid, data} hoặc {valid, error}", "Không ghi file khi valid = false."),
        ],
        [2100, 4000, 3260],
        font_size=9.0,
    )

    page_break(doc)

    # Page 6
    add_heading(doc, "5. Cách tổ chức source code", 1)
    add_code_block(
        doc,
        "K3-Day9-Multi-Agent-A2A/\n"
        "├── main.py                     # batch runner cho 50 case\n"
        "├── agents/\n"
        "│   ├── base_agent.py           # interface process()\n"
        "│   ├── coordinator_agent.py    # orchestration + assembly\n"
        "│   ├── order_seller_agent.py   # order/item/seller\n"
        "│   ├── payment_agent.py        # financial reconciliation\n"
        "│   ├── delivery_agent.py       # delivery timing\n"
        "│   ├── policy_agent.py         # EC_POLICY_V1\n"
        "│   ├── evidence_agent.py       # evidence selection\n"
        "│   └── verifier_agent.py       # validation\n"
        "├── core/\n"
        "│   ├── data_loader.py          # CSV access layer\n"
        "│   ├── schemas.py              # Pydantic output schema\n"
        "│   ├── tracer.py               # JSONL audit log\n"
        "│   └── llm_client.py           # OpenAI client (hiện chưa được gọi)\n"
        "├── data/                        # Olist CSV\n"
        "├── input/                       # 50 case JSON\n"
        "├── output/                      # 50 result JSON\n"
        "└── logging/                     # trace.jsonl, metadata.json"
    )
    add_heading(doc, "5.1. Nguyên tắc thiết kế", 2)
    principles = [
        ("Separation of concerns", "Mỗi domain có agent riêng; Policy không kiêm tính tổng payment hay so sánh timestamp."),
        ("Determinism", "Item, payment và seller được sắp xếp; luật if/elif cố định; tiền làm tròn trước khi xuất."),
        ("Evidence-first", "Evidence ID phải dựng trực tiếp từ dữ liệu hoặc mã policy hợp lệ."),
        ("Fail closed", "Case không khớp luật hoặc sai schema sẽ raise lỗi, không âm thầm tạo output không chắc chắn."),
        ("Auditability", "Mọi kết quả trung gian và final output được lưu trong một record trace theo case."),
    ]
    add_table(doc, ["Nguyên tắc", "Biểu hiện trong mã nguồn"], principles, [2700, 6660], font_size=9.3)

    page_break(doc)

    # Page 7
    add_heading(doc, "6. Luật nghiệp vụ EC_POLICY_V1", 1)
    add_para(doc, "PolicyAgent duyệt luật theo đúng thứ tự dưới đây. Thứ tự là một phần của nghiệp vụ: khi case thỏa nhiều dấu hiệu, luật xuất hiện trước được ưu tiên.")
    policy_rows = [
        ("1", "canceled_order_paid", "status=canceled, payment>0", "platform", "Tổng payment", "issue_full_refund"),
        ("2", "unavailable_order_paid", "status=unavailable, payment>0", "platform", "Tổng payment", "issue_full_refund"),
        ("3", "late_delivery_seller", "giao trễ + carrier nhận sau limit", "seller", "Tổng freight", "refund_freight"),
        ("4", "late_delivery_logistics", "giao trễ + carrier nhận đúng limit", "logistics", "Tổng freight", "refund_freight"),
        ("5", "valid_split_payment", "≥2 payment + tổng tiền khớp", "không có", "0", "explain_valid_split_payment"),
        ("6", "unsupported_late_claim", "không giao trễ + tổng tiền khớp", "không có", "0", "reject_late_refund"),
    ]
    add_table(doc, ["#", "Primary issue", "Điều kiện rút gọn", "Chịu trách nhiệm", "Refund", "Action"], policy_rows, [350, 1900, 2450, 1500, 1250, 1910], font_size=8.2)
    add_heading(doc, "6.1. Cách xác định nguyên nhân giao trễ", 2)
    add_para(doc, "DeliveryAgent trước hết kiểm tra order_delivered_customer_date > order_estimated_delivery_date. Chỉ khi đơn thực sự trễ mới xét thời điểm carrier nhận hàng:")
    add_bullet(doc, "Nếu order_delivered_carrier_date > shipping_limit_date của item: seller bàn giao trễ → trách nhiệm seller.")
    add_bullet(doc, "Nếu carrier nhận không trễ nhưng khách nhận sau estimated date: trách nhiệm logistics provider.")
    add_heading(doc, "6.2. Đối soát tài chính", 2)
    add_para(doc, "PaymentAgent tính payment_total = Σ payment_value, item_total = Σ price và freight_total = Σ freight_value. Tổng thanh toán được coi là khớp khi |payment_total − (item_total + freight_total)| ≤ 0,10 BRL.")
    add_callout(doc, "Case status", "recommended_refund_brl > 0 → action_required; ngược lại → no_action. Verifier kiểm tra lại quan hệ này trước khi ghi file.", tone="blue")

    page_break(doc)

    # Page 8
    add_heading(doc, "7. Output, evidence và kiểm định", 1)
    add_heading(doc, "7.1. Cấu trúc output", 2)
    add_table(
        doc,
        ["Nhóm trường", "Nội dung"],
        [
            ("assessment", "primary_issue, case_status, confidence"),
            ("affected_entities", "order_ids, item_ids, seller_ids, payment_ids"),
            ("root_cause_analysis", "ranked_causes và responsible_parties"),
            ("evidence_ids", "Order/item/payment/seller/policy IDs có thể kiểm chứng"),
            ("financial_resolution", "Currency, item/freight/payment total và recommended refund"),
            ("resolution_actions", "Hành động xử lý theo policy"),
        ],
        [2500, 6860],
        font_size=9.4,
    )
    add_heading(doc, "7.2. Affected entities khác evidence", 2)
    add_para(doc, "Affected entities là toàn bộ đối tượng liên quan đến case. Evidence là tập bản ghi cần để chứng minh kết luận và các tổng tiền được báo cáo. Vì vậy một seller có thể nằm trong seller_ids nhưng chỉ trở thành seller:<id> trong evidence khi seller là bên chịu trách nhiệm.")
    add_heading(doc, "7.3. EvidenceAgent", 2)
    for text in (
        "Luôn thêm order:<order_id>.",
        "Thêm các item row khi chúng tồn tại để chứng minh item/freight total.",
        "Thêm mọi payment row tham gia payment_total và đối soát.",
        "Chỉ thêm seller record khi responsible_party_type = seller.",
        "Thêm policy:<root_cause_code> để chỉ rõ luật đã áp dụng.",
    ):
        add_bullet(doc, text)
    add_heading(doc, "7.4. VerifierAgent và trace", 2)
    add_para(doc, "Verifier cắt list theo giới hạn đề bài (5 entity ID mỗi nhóm, 10 evidence, 3 cause/party, 5 action), làm tròn tiền 2 chữ số, chặn confidence trong [0,1], đồng bộ case_status và validate Pydantic. Tracer lưu output từng agent cùng final output thành một dòng logging/trace.jsonl cho mỗi case.")

    page_break(doc)

    # Page 9
    add_heading(doc, "8. Cách chạy lại project", 1)
    add_heading(doc, "8.1. Cài môi trường", 2)
    add_code_block(
        doc,
        "cd C:\\Users\\Admin\\Downloads\\K3-Day9-Multi-Agent-A2A\n"
        ".\\.venv\\Scripts\\Activate.ps1\n"
        "python -m pip install -r requirements.txt"
    )
    add_para(doc, "Nếu PowerShell chặn Activate.ps1, có thể gọi trực tiếp .venv\\Scripts\\python.exe cho các lệnh Python.")
    add_heading(doc, "8.2. Chạy batch", 2)
    add_code_block(doc, "python main.py")
    add_heading(doc, "8.3. Kiểm tra artifact", 2)
    add_code_block(
        doc,
        "(Get-ChildItem output\\EC_*.json).Count       # mong đợi: 50\n"
        "(Get-Content logging\\trace.jsonl).Count      # mong đợi: 50\n"
        "Compress-Archive -Path output\\* -DestinationPath output.zip -Force"
    )
    add_callout(
        doc,
        "Vì sao chương trình chạy nhanh?",
        "CSV được nạp một lần vào Pandas; mỗi case chỉ lọc dữ liệu và chạy các phép tính/if-else trong bộ nhớ. "
        "Mặc dù LLMClient được khởi tạo, các agent hiện không gọi chat() hoặc get_json(), nên không có độ trễ mạng hay thời gian sinh token.",
        tone="green",
    )
    add_heading(doc, "8.4. Ý nghĩa confidence = 1.0", 2)
    add_para(doc, "Trong phiên bản hiện tại, confidence = 1.0 không phải xác suất do mô hình học máy dự đoán. Nó là cờ quy ước: case khớp đầy đủ một luật EC_POLICY_V1 xác định từ các trường CSV được cung cấp. Vì vậy confidence phản ánh độ chắc chắn của rule match trong phạm vi dữ liệu, không khẳng định dữ liệu nguồn ngoài đời là hoàn hảo.")

    page_break(doc)

    # Page 10
    add_heading(doc, "9. Đánh giá thiết kế và giới hạn hiện tại", 1)
    add_heading(doc, "9.1. Điểm mạnh", 2)
    strengths = [
        "Phân tách đúng các domain order/seller, payment, delivery, policy, evidence và verification.",
        "Quyết định ổn định, có thể tái hiện và dễ so sánh với ground truth của bài chấm.",
        "Các ID và danh sách được chuẩn hóa thứ tự, hạn chế sai khác không cần thiết giữa các lượt chạy.",
        "Có validation cuối luồng và trace đủ để giải thích từng case.",
        "Không phụ thuộc kết quả sinh tự do của LLM cho số tiền hoặc sự kiện quan trọng.",
    ]
    for text in strengths:
        add_bullet(doc, text)
    add_heading(doc, "9.2. Giới hạn cần trình bày trung thực", 2)
    limitations = [
        "Bảy agent là bảy lớp xử lý chuyên trách, không phải bảy phiên LLM độc lập. LLMClient tồn tại nhưng hiện không được gọi trong pipeline.",
        "Ba specialist chạy tuần tự; kiến trúc cho phép song song hóa nhưng mã hiện tại chưa dùng concurrency.",
        "Policy chỉ xử lý sáu nhóm case trong EC_POLICY_V1; trường hợp ngoài policy sẽ raise lỗi.",
        "Pydantic schema kiểm tra kiểu/cấu trúc nhưng chưa kiểm tra evidence ID thực sự tồn tại bằng một lookup riêng ở Verifier.",
        "metadata.json ghi gpt-4o-mini ~8B, nhưng con số tham số này không có xác nhận chính thức trong source của project. Không nên dùng metadata đó làm bằng chứng duy nhất cho yêu cầu ≤10B.",
    ]
    for text in limitations:
        add_bullet(doc, text)
    add_callout(
        doc,
        "Cách diễn đạt khi báo cáo",
        "Hệ thống hiện là multi-agent deterministic orchestration: agent phân chia vai trò và handoff bằng object Python; quyết định nghiệp vụ chạy bằng rule trên dữ liệu kiểm chứng. "
        "OpenAI SDK là thành phần dự phòng, chưa tham gia inference ở lượt chạy output hiện tại.",
        tone="amber",
    )
    add_heading(doc, "9.3. Hướng cải tiến", 2)
    add_table(
        doc,
        ["Ưu tiên", "Cải tiến", "Lợi ích"],
        [
            ("1", "Thêm EvidenceValidator tra ngược ID vào DataLoader", "Giảm false-positive evidence."),
            ("2", "Chạy 3 specialist song song", "Giảm latency khi logic hoặc nguồn dữ liệu lớn hơn."),
            ("3", "Thêm unit test cho 6 rule và boundary 0,10 BRL", "Ngăn regression khi sửa pipeline."),
            ("4", "Nếu bắt buộc dùng LLM, chọn model có tài liệu ≤10B", "Chứng minh tuân thủ đề bài rõ ràng."),
        ],
        [850, 4450, 4060],
        font_size=9.2,
    )

    page_break(doc)

    # Page 11
    add_heading(doc, "10. Kịch bản trình bày ngắn", 1)
    add_para(doc, "Có thể trình bày dự án trong khoảng 4–5 phút theo thứ tự sau:")
    presentation = [
        ("Bài toán (30 giây)", "Hệ thống xử lý 50 khiếu nại Olist, phải đối chiếu nhiều bảng và xuất kết luận có bằng chứng."),
        ("Kiến trúc (60 giây)", "Coordinator điều phối; ba specialist phân tích order/payment/delivery; Policy quyết định; Evidence chọn bằng chứng; Verifier chặn output sai."),
        ("Luồng dữ liệu (60 giây)", "Input → lọc CSV theo order_id → handoff kết quả chuẩn hóa → EC_POLICY_V1 → output + trace."),
        ("Quyết định nghiệp vụ (60 giây)", "Giải thích một ví dụ giao trễ: seller giao carrier sau shipping limit thì seller chịu trách nhiệm; ngược lại logistics chịu trách nhiệm."),
        ("Tính đúng và audit (45 giây)", "Nêu evidence ID, giới hạn schema, làm tròn tiền và Pydantic validation."),
        ("Kết luận (30 giây)", "Thiết kế ưu tiên tính kiểm chứng và tái lập; pipeline rule-based nên chạy nhanh và ổn định."),
    ]
    add_table(doc, ["Phần", "Nội dung cần nói"], presentation, [2300, 7060], font_size=9.6)
    add_heading(doc, "10.1. Ba câu hỏi thường gặp", 2)
    qa = [
        ("Tại sao gọi là multi-agent khi không gọi LLM?", "Vì hệ thống phân rã vai trò, contract và handoff giữa nhiều tác nhân phần mềm; LLM không phải điều kiện bắt buộc của khái niệm agent trong thiết kế này."),
        ("Tại sao confidence bằng 1.0?", "Vì đây là độ chắc chắn quy ước của rule match xác định, không phải probability do model dự đoán."),
        ("Tại sao output chạy rất nhanh?", "Dữ liệu đã nằm trong RAM, quyết định là phép tính và if/else, không gọi API mạng trong từng case."),
    ]
    add_table(doc, ["Câu hỏi", "Câu trả lời ngắn"], qa, [3500, 5860], font_size=9.4)
    add_heading(doc, "11. Kết luận", 1)
    add_para(doc, "Project đáp ứng bài toán bằng kiến trúc 7 agent có vai trò và handoff rõ. Điểm cốt lõi của thiết kế là tách fact extraction, policy decision, evidence selection và validation thành các lớp độc lập. Cách làm này phù hợp với bài chấm tự động vì kết quả có tính xác định, có thể audit và bám sát EC_POLICY_V1.")
    add_callout(doc, "Thông điệp chốt", "Đây là một pipeline multi-agent hướng bằng chứng: dữ liệu quyết định kết luận, policy quyết định hành động và verifier bảo vệ chất lượng output.", tone="green")

    add_header_footer(doc)
    doc.core_properties.title = "Báo cáo thiết kế và luồng hoạt động hệ thống Multi-Agent Olist"
    doc.core_properties.subject = "K3 Day 09 - Multi-Agent A2A"
    doc.core_properties.author = "Vũ Việt Anh"
    doc.core_properties.keywords = "multi-agent, Olist, dispute resolution, EC_POLICY_V1"
    doc.save(REPORT_PATH)
    print(REPORT_PATH)


if __name__ == "__main__":
    build_report()
